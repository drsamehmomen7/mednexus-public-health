"""
Extracts plain text from an uploaded document, so the SAME extraction
pipeline that already works on pasted text can run on uploaded files too
— this module's only job is "bytes in, text out", nothing report-type
specific.

DOCX (python-docx), TXT and PDF (pypdf, text layer only — no OCR). CSV is
mentioned in the frontend's dropzone hint as planned — add a branch here
per format as each one is actually built, rather than promising formats
this doesn't handle yet.
"""

import logging
import re
from io import BytesIO
from typing import Optional

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PasswordType, PdfReader
from pypdf.errors import DependencyError

logger = logging.getLogger(__name__)

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PAGES = 30
# DOCX and TXT have no page objects to count, so their length is judged by an estimate of ~3,000 characters per page.
CHARS_PER_PAGE_ESTIMATE = 3000

NOT_A_PDF_MESSAGE = (
    "That file has a .pdf extension but isn't a valid PDF (it may be empty, "
    "damaged, or a different file type that was renamed)."
)
CORRUPT_PDF_MESSAGE = (
    "This PDF couldn't be read. It may be damaged or use a feature MedNexus "
    "doesn't support yet. Try re-saving or re-exporting it as a new PDF."
)
ENCRYPTED_PDF_MESSAGE = (
    "This PDF is password-protected (encrypted), so MedNexus can't read it. "
    "Upload an unprotected copy."
)
NO_TEXT_PDF_MESSAGE = (
    "No readable text found in that document. This PDF looks like a scan or "
    "image-only file (no selectable text), and OCR isn't supported yet. Upload "
    "a text-based PDF or paste the report text instead."
)
CORRUPT_DOCX_MESSAGE = (
    "This Word file couldn't be read. It may be damaged or not a real .docx "
    "document. Try re-saving it as a new .docx."
)


class UnsupportedDocumentType(Exception):
    """Raised when the uploaded file's extension isn't handled yet."""


class DocumentTooLarge(Exception):
    """Raised when the upload exceeds MAX_UPLOAD_BYTES. The message is user-facing."""


class UnreadableDocument(Exception):
    """Raised when a supported file can't be read (damaged, encrypted, no text, too long). The message is user-facing."""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Return the document's text content in TRUE reading order — paragraphs
    and tables interleaved exactly as they appear in the document body,
    not all paragraphs followed by all tables.

    This matters for any real form that has a letterhead paragraph, then
    a field table, then a signature-line footer paragraph (the common
    shape of an official reporting form): grouping by element type would
    put the footer text ahead of the table content it actually follows,
    which is disorienting to review even though it happened not to
    break extraction in testing. python-docx doesn't expose body order
    directly, so this walks the underlying XML body children instead.
    """
    document = Document(BytesIO(file_bytes))
    parts = []

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = Paragraph(child, document)
            if paragraph.text.strip():
                parts.append(paragraph.text)

        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" ".join(cells))

    return "\n".join(parts)


# Keys are code points because most of these characters are invisible in source. Only what would break matching is normalized.
_PDF_CHAR_MAP = str.maketrans(
    {
        0xFB00: "ff", 0xFB01: "fi", 0xFB02: "fl", 0xFB03: "ffi", 0xFB04: "ffl", 0xFB05: "st", 0xFB06: "st",  # ligatures: "influenza" with an fl ligature would miss the gazetteer
        **dict.fromkeys([0x00A0, 0x202F, 0x205F, 0x3000, *range(0x2000, 0x200B)], " "),  # no-break and other exotic spaces
        **dict.fromkeys([0x00AD, 0x200B, 0x200C, 0x200D, 0x2060, 0xFEFF], None),  # soft hyphen, zero-width characters, BOM
        0x2028: "\n",  # line separator
        0x2029: "\n",  # paragraph separator
        0xF0B7: chr(0x2022),  # bullet as encoded by the Symbol font
    }
)


def normalize_pdf_text(text: str) -> str:
    """Character clean-up plus whitespace tidying. Deliberately does NOT unwrap lines, dehyphenate or strip repeated headers."""
    text = text.translate(_PDF_CHAR_MAP).replace("\r\n", "\n").replace("\r", "\n").replace("\f", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Return the PDF's text layer, pages separated by a blank line. No OCR: a
    scan (no text layer) raises UnreadableDocument, as do encrypted, damaged
    and over-long files — every message here is safe to show the user.
    """
    if b"%PDF-" not in file_bytes[:1024]:
        raise UnreadableDocument(NOT_A_PDF_MESSAGE)

    try:
        reader = PdfReader(BytesIO(file_bytes))
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:  # owner-only restrictions open with an empty user password
            raise UnreadableDocument(ENCRYPTED_PDF_MESSAGE)

        page_count = len(reader.pages)
        if page_count == 0:
            raise UnreadableDocument(CORRUPT_PDF_MESSAGE)
        if page_count > MAX_PAGES:
            raise UnreadableDocument(
                f"That PDF has {page_count} pages; the limit is {MAX_PAGES}. "
                "Reports are expected to be short."
            )
        pages = [normalize_pdf_text(page.extract_text() or "") for page in reader.pages]
    except UnreadableDocument:
        raise
    except DependencyError as exc:
        # AES-encrypted and no crypto package installed (by design). pypdf raises this from the constructor (AES-256) or later from extract_text()
        # (AES-128 with only an owner password); AES-128 with a user password is caught earlier, where decrypt("") returns NOT_DECRYPTED.
        if "AES" in str(exc):
            raise UnreadableDocument(ENCRYPTED_PDF_MESSAGE) from None
        logger.warning("PDF could not be read (DependencyError)")
        raise UnreadableDocument(CORRUPT_PDF_MESSAGE) from exc
    except Exception as exc:
        # Third-party parser on untrusted input: whatever it throws, the honest answer is "can't read this file".
        logger.warning("PDF could not be read (%s)", type(exc).__name__)
        raise UnreadableDocument(CORRUPT_PDF_MESSAGE) from exc

    text = "\n\n".join(page for page in pages if page)
    if not text.strip():
        raise UnreadableDocument(NO_TEXT_PDF_MESSAGE)
    return text


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Dispatch by file extension. Raises UnsupportedDocumentType for
    anything not yet handled, so the caller can return a clear error
    instead of silently returning empty text. DocumentTooLarge and
    UnreadableDocument carry user-facing messages for the other ways a
    file can fail; the size and page limits apply to every format.
    """
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise DocumentTooLarge(f"That file is larger than the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit.")

    lowered = (filename or "").lower()
    if lowered.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)  # page limit is enforced there, on real page objects
    if lowered.endswith(".docx"):
        try:
            text = extract_text_from_docx(file_bytes)
        except Exception as exc:
            logger.warning("DOCX could not be read (%s)", type(exc).__name__)
            raise UnreadableDocument(CORRUPT_DOCX_MESSAGE) from exc
    elif lowered.endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocumentType(
            f"'{filename}' isn't a supported format yet. DOCX, TXT and PDF are handled; "
            "CSV isn't built yet."
        )

    approx_pages = -(-len(text) // CHARS_PER_PAGE_ESTIMATE)
    if approx_pages > MAX_PAGES:
        raise UnreadableDocument(
            f"That document is about {approx_pages} pages long; the limit is {MAX_PAGES}. "
            "Reports are expected to be short."
        )
    return text
